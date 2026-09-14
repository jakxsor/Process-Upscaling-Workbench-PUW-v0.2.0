"""Native PowerPoint export for the editable flowsheet view.

The generated deck intentionally uses ordinary PowerPoint shapes, connector
segments, and text boxes instead of embedding the SVG as an image. That keeps
the slide editable after export.
"""

from __future__ import annotations

import io
import re
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Emu, Pt


EMU_PER_INCH = 914400


def _clean(value: Any, default: str = "") -> str:
    text = str(value or "").strip()
    return re.sub(r"\s+", " ", text) if text else default


def _color(value: str) -> str:
    return re.sub(r"[^0-9a-fA-F]", "", value or "")[:6].upper() or "172027"


def _float(value: Any, default: float = 0) -> float:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return default
    return result if result == result else default


def _stream_label(stream: dict[str, Any]) -> str:
    name = _clean(stream.get("name"), "stream")
    qty = _clean(stream.get("quantity"))
    unit = _clean(stream.get("unit"))
    return f"{name}{', ' + qty + (' ' + unit if unit else '') if qty else ''}"


def _mass_kg(stream: dict[str, Any]) -> float:
    qty = _clean(stream.get("quantity")).replace(",", ".")
    match = re.search(r"[-+]?\d+(?:\.\d+)?", qty)
    if not match:
        return 0
    value = float(match.group(0))
    unit = _clean(stream.get("unit")).lower()
    if unit in ("kg", "kilogram", "kilograms"):
        return value
    if unit in ("g", "gram", "grams"):
        return value / 1000
    if unit in ("mg", "milligram", "milligrams"):
        return value / 1_000_000
    return 0


def _link_mass_kg(link: dict[str, Any]) -> float:
    return sum(_mass_kg(stream) for stream in link.get("directStreams") or [] if isinstance(stream, dict))


def _clip_words(value: Any, max_chars: int, min_chars: int = 8) -> str:
    text = _clean(value)
    if len(text) <= max_chars:
        return text
    if max_chars < min_chars + 3:
        return ""
    prefix = text[: max_chars - 3]
    cut = max(prefix.rfind(" "), prefix.rfind("/"), prefix.rfind("-"))
    clipped = prefix[:cut].strip() if cut >= min_chars else prefix.strip()
    return f"{clipped}..." if clipped else ""


def _compact_link_label(link: dict[str, Any], kind: str = "process") -> str:
    streams = [item for item in link.get("directStreams") or [] if isinstance(item, dict)]
    if not streams:
        return kind
    stream = streams[0]
    name = _clean(stream.get("name"), "stream")
    qty = _clean(stream.get("quantity"))
    unit = _clean(stream.get("unit"))
    full = f"{name}{', ' + qty + (' ' + unit if unit else '') if qty else ''}"
    return _clip_words(full, 28) or _clip_words(name, 24) or kind


def _category_label(category: Any) -> str:
    labels = {
        "reactor": "Reactor",
        "separation": "Separation",
        "utility": "Utility",
        "storage": "Storage",
        "waste": "Waste",
    }
    return labels.get(_clean(category), "Unit")


class _ShapeWriter:
    def __init__(self, slide: Any, scale: float, x_offset: float, y_offset: float) -> None:
        self.slide = slide
        self.scale = scale
        self.x_offset = x_offset
        self.y_offset = y_offset

    def _emu(self, value: float) -> int:
        return int(round(value))

    def map_x(self, px: Any) -> int:
        return self._emu(self.x_offset + _float(px) * self.scale)

    def map_y(self, px: Any) -> int:
        return self._emu(self.y_offset + _float(px) * self.scale)

    def map_len(self, px: Any) -> int:
        return max(1, self._emu(_float(px) * self.scale))

    def rect(self, x: Any, y: Any, w: Any, h: Any, *, fill: str, line: str, radius: bool = False, name: str = "Shape") -> None:
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        shape = self.slide.shapes.add_shape(
            kind, Emu(self.map_x(x)), Emu(self.map_y(y)), Emu(self.map_len(w)), Emu(self.map_len(h))
        )
        shape.name = _clean(name, "Shape")
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(_color(fill))
        shape.line.color.rgb = RGBColor.from_string(_color(line))
        shape.line.width = Pt(1)

    def textbox(self, x: Any, y: Any, w: Any, h: Any, text: str, *, size: float = 10, color: str = "172027", bold: bool = False, name: str = "Text", max_lines: int = 4) -> None:
        lines = _wrap(text, max(10, int(_float(w) / max(size * 0.45, 4))), max_lines)
        shape = self.slide.shapes.add_textbox(
            Emu(self.map_x(x)), Emu(self.map_y(y)), Emu(self.map_len(w)), Emu(self.map_len(h))
        )
        shape.name = _clean(name, "Text")
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = frame.margin_right = Emu(0)
        frame.margin_top = frame.margin_bottom = Emu(0)
        for index, line_text in enumerate(lines or [""]):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = line_text
            paragraph.font.name = "Aptos"
            paragraph.font.size = Pt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(_color(color))
            paragraph.space_before = paragraph.space_after = Pt(0)

    def label(self, x: Any, y: Any, w: Any, text: str, *, color: str = "172027", name: str = "Label") -> None:
        if not _clean(text):
            return
        self.rect(x, y, w, 18, fill="FFFFFF", line="D6E0E5", radius=True, name=f"{name} background")
        self.textbox(_float(x) + 6, _float(y) + 3, max(1, _float(w) - 12), 12, text, size=8.3, color=color, bold=True, name=name, max_lines=1)

    def line(self, x1: Any, y1: Any, x2: Any, y2: Any, *, color: str = "172027", width: float = 2.2, arrow: bool = False, dash: bool = False) -> None:
        x1e, y1e = self.map_x(x1), self.map_y(y1)
        x2e, y2e = self.map_x(x2), self.map_y(y2)
        connector = self.slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, Emu(x1e), Emu(y1e), Emu(x2e), Emu(y2e)
        )
        connector.name = "Connector"
        connector.line.color.rgb = RGBColor.from_string(_color(color))
        connector.line.width = Pt(max(0.5, width))
        if dash:
            connector.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        if arrow:
            connector.line._get_or_add_ln().append(
                parse_xml(f'<a:tailEnd {nsdecls("a")} type="triangle"/>')
            )


def _wrap(text: Any, width: int, max_lines: int = 2) -> list[str]:
    words = _clean(text).split()
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if current and len(candidate) > width:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines[:max_lines]


def _center(box: dict[str, Any]) -> tuple[float, float]:
    return _float(box.get("x")) + _float(box.get("w")) / 2, _float(box.get("y")) + _float(box.get("h")) / 2


def _port(from_box: dict[str, Any], to_box: dict[str, Any], offset: float = 12) -> tuple[float, float]:
    from_x, from_y = _center(from_box)
    to_x, to_y = _center(to_box)
    dx, dy = to_x - from_x, to_y - from_y
    x, y, w, h = (_float(from_box.get("x")), _float(from_box.get("y")), _float(from_box.get("w")), _float(from_box.get("h")))
    if abs(dx) >= abs(dy):
        return (x + w + offset, from_y) if dx >= 0 else (x - offset, from_y)
    return (from_x, y + h + offset) if dy >= 0 else (from_x, y - offset)


def _route(from_box: dict[str, Any], to_box: dict[str, Any]) -> list[tuple[float, float]]:
    start = _port(from_box, to_box, 10)
    end = _port(to_box, from_box, 14)
    if abs(start[1] - end[1]) < 3 or abs(start[0] - end[0]) < 3:
        return [start, end]
    if abs(start[0] - end[0]) >= abs(start[1] - end[1]):
        mid_x = (start[0] + end[0]) / 2
        return [start, (mid_x, start[1]), (mid_x, end[1]), end]
    mid_y = (start[1] + end[1]) / 2
    return [start, (start[0], mid_y), (end[0], mid_y), end]


def _link_points(link: dict[str, Any], src: dict[str, Any], dst: dict[str, Any]) -> list[tuple[float, float]]:
    """The route the SVG drew for this link, if the export carried it; else a local route.

    The slide used to re-route every arrow with its own rules, so it never matched the
    drawing on screen. The check fixture carries no points, so the fallback stays.
    """
    raw = link.get("points")
    points: list[tuple[float, float]] = []
    if isinstance(raw, list):
        for item in raw:
            if isinstance(item, dict) and "x" in item and "y" in item:
                points.append((_float(item.get("x")), _float(item.get("y"))))
    if len(points) >= 2:
        return points
    return _route(src, dst)


def _link_label(link: dict[str, Any], kind: str = "process") -> str:
    """Prefer the label the SVG printed (stream number, total mass); else the local compact one."""
    label = _clean(link.get("label"))
    return label if label else _compact_link_label(link, kind)


def _draw_polyline(writer: _ShapeWriter, points: list[tuple[float, float]], *, color: str, width: float, arrow: bool = True, dash: bool = False) -> None:
    clean_points = [point for index, point in enumerate(points) if index == 0 or point != points[index - 1]]
    for index in range(len(clean_points) - 1):
        x1, y1 = clean_points[index]
        x2, y2 = clean_points[index + 1]
        writer.line(x1, y1, x2, y2, color=color, width=width, arrow=arrow and index == len(clean_points) - 2, dash=dash)


def _label_anchor(points: list[tuple[float, float]]) -> tuple[float, float, float] | None:
    best: tuple[float, float, float] | None = None
    for index in range(len(points) - 1):
        a = points[index]
        b = points[index + 1]
        if abs(a[1] - b[1]) > 1:
            continue
        length = abs(b[0] - a[0])
        if best is None or length > best[2]:
            best = (min(a[0], b[0]) + 6, a[1] - 23, length - 12)
    if best is None or best[2] < 88:
        return None
    return best


def _draw_link_label(writer: _ShapeWriter, points: list[tuple[float, float]], label: str, color: str) -> None:
    anchor = _label_anchor(points)
    if not anchor:
        return
    x, y, available = anchor
    compact = _clip_words(label, max(10, int(available / 5.5)), 8)
    if not compact:
        return
    writer.label(x, y, min(available, max(74, len(compact) * 5.6 + 18)), compact, color=color, name="Stream label")


def _draw_stream_table(writer: _ShapeWriter, flowsheet: dict[str, Any]) -> None:
    """The stream table from the drawing as a real, editable PowerPoint table."""
    rows = [row for row in flowsheet.get("streamTable") or [] if isinstance(row, dict)]
    top = _float(flowsheet.get("tableTop"), 0)
    if not rows or top <= 0:
        return
    width_px = max(600, _float(flowsheet.get("width"), 1200) - 72)
    row_h = 15
    left = writer.map_x(36)
    table_top = writer.map_y(top + 4)
    table_w = writer.map_len(width_px)
    table_h = writer.map_len(row_h * (len(rows) + 1) + 6)
    shape = writer.slide.shapes.add_table(len(rows) + 1, 6, Emu(left), Emu(table_top), Emu(table_w), Emu(table_h))
    shape.name = "Stream table"
    table = shape.table
    widths = [0.05, 0.11, 0.08, 0.13, 0.07, 0.56]
    for index, fraction in enumerate(widths):
        table.columns[index].width = Emu(int(table_w * fraction))
    headers = ["No.", "From > To", "Type", "Total", "Phase", "Composition (largest first)"]

    def write_cell(cell: Any, text: str, *, bold: bool = False) -> None:
        cell.text = text
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(7.5)
                run.font.bold = bold
                run.font.color.rgb = RGBColor.from_string("172027")
        cell.margin_left = Emu(int(0.04 * EMU_PER_INCH))
        cell.margin_right = Emu(int(0.04 * EMU_PER_INCH))
        cell.margin_top = Emu(int(0.01 * EMU_PER_INCH))
        cell.margin_bottom = Emu(int(0.01 * EMU_PER_INCH))

    for column, header in enumerate(headers):
        write_cell(table.cell(0, column), header, bold=True)
    for row_index, row in enumerate(rows, start=1):
        total_kg = row.get("totalKg")
        unknown = int(_float(row.get("unknown"), 0))
        count = int(_float(row.get("count"), 0))
        if isinstance(total_kg, (int, float)) and total_kg == total_kg:
            mass = f"{total_kg:.3g} kg" if total_kg < 100 else f"{total_kg:.0f} kg"
            if unknown:
                mass += f" (+{unknown} n.q.)"
        else:
            mass = f"n.q. ({count} streams)" if count else "no declared streams"
        composition = row.get("composition") if isinstance(row.get("composition"), list) else []
        values = [
            _clean(row.get("tag")),
            f"{_clean(row.get('from'))} > {_clean(row.get('to'))}",
            _clean(row.get("kind")),
            mass,
            _clean(row.get("phase"), "-"),
            "; ".join(_clean(item) for item in composition),
        ]
        for column, value in enumerate(values):
            write_cell(table.cell(row_index, column), value, bold=column == 0)


def _draw_title_block(writer: _ShapeWriter, flowsheet: dict[str, Any], diagram_w: float, diagram_h: float) -> None:
    block = flowsheet.get("titleBlock")
    if not isinstance(block, dict):
        return
    x = max(620, diagram_w - 470)
    y = diagram_h - 92
    writer.rect(x, y, 440, 56, fill="FFFFFF", line="172027", name="Title block")
    writer.textbox(x + 8, y + 4, 110, 18, "DRAWING", size=8, bold=True)
    missing = int(_float(block.get("missing"), 0))
    writer.textbox(x + 120, y + 4, 310, 18, f"{_clean(block.get('drawing'), 'PFD-01')} - {_clean(block.get('date'))} - rev. {_clean(block.get('revision'), 'draft')}", size=8)
    writer.textbox(x + 8, y + 30, 110, 18, "BASIS", size=8, bold=True)
    basis = _clean(block.get("basis"), "lab batch")
    writer.textbox(x + 120, y + 30, 310, 18, f"{basis}{f' - {missing} n.q.' if missing else ''}", size=8)


def _draw_legend(writer: _ShapeWriter, diagram_w: float, diagram_h: float, category_styles: dict[str, tuple[str, str]], legend_y: float | None = None) -> None:
    x = 36
    y = max(96, diagram_h - 78) if legend_y is None else legend_y
    writer.rect(x - 10, y - 18, min(790, diagram_w - 52), 58, fill="FFFFFF", line="D6E0E5", radius=True, name="Legend Panel")
    writer.textbox(x, y - 12, 72, 14, "Legend", size=9.5, color="657480", bold=True, name="Legend title")

    line_items = [
        ("Process", "172027", False),
        ("Recycle", "25834A", True),
        ("Waste", "965D00", False),
        ("Vent/VOC", "657480", True),
    ]
    cursor_x = x + 78
    for label, color, dash in line_items:
        writer.line(cursor_x, y - 4, cursor_x + 28, y - 4, color=color, width=2.2, arrow=True, dash=dash)
        writer.textbox(cursor_x + 36, y - 11, 70, 14, label, size=8.2, color=color, bold=True, name=f"Legend {label}")
        cursor_x += 114

    category_items = ["reactor", "separation", "utility", "storage", "waste"]
    cursor_x = x + 78
    for category in category_items:
        fill, line = category_styles[category]
        writer.rect(cursor_x, y + 17, 13, 10, fill=fill, line=line, radius=True, name=f"Legend {_category_label(category)} swatch")
        writer.textbox(cursor_x + 18, y + 15, 72, 14, _category_label(category), size=8.0, color="40515D", bold=True, name=f"Legend {_category_label(category)}")
        cursor_x += 104


def render_flowsheet_pptx(payload: dict[str, Any]) -> bytes:
    flowsheet = payload.get("flowsheet") if isinstance(payload, dict) else None
    if not isinstance(flowsheet, dict):
        raise ValueError("Missing flowsheet export model.")

    groups = [item for item in flowsheet.get("groups") or [] if isinstance(item, dict)]
    if not groups:
        raise ValueError("No grouped operations available for PowerPoint export.")

    diagram_w = max(900, _float(flowsheet.get("width"), 1880))
    diagram_h = max(620, _float(flowsheet.get("height"), 760))
    slide_w = int(min(22.0, max(13.333, diagram_w / 96)) * EMU_PER_INCH)
    slide_h = int(min(14.0, max(7.5, diagram_h / 96)) * EMU_PER_INCH)
    margin = int(0.28 * EMU_PER_INCH)
    scale = min((slide_w - margin * 2) / diagram_w, (slide_h - margin * 2) / diagram_h)
    presentation = Presentation()
    presentation.slide_width = slide_w
    presentation.slide_height = slide_h
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    writer = _ShapeWriter(slide, scale, (slide_w - diagram_w * scale) / 2, (slide_h - diagram_h * scale) / 2)

    category_styles = {
        "reactor": ("FFF8F7", "CF4B42"),
        "separation": ("F5F9FF", "1671C2"),
        "utility": ("FFFAF0", "B97916"),
        "storage": ("F1FBFA", "0F766E"),
        "waste": ("FAF8F5", "8B7057"),
    }

    writer.rect(0, 0, diagram_w, diagram_h, fill="FFFFFF", line="172027", name="Drawing Border")
    writer.textbox(30, 28, 620, 32, _clean(flowsheet.get("title"), "Generated Process Flowsheet"), size=18, bold=True)
    writer.textbox(30, 62, 980, 28, _clean(flowsheet.get("basis"), "Editable PowerPoint export from declared process links."), size=10.5, color="657480")

    by_id = {_clean(group.get("id")): group for group in groups}
    max_kg = max(0.01, *[_float(group.get("totalOutputKg")) for group in groups])

    def stroke_width(link: dict[str, Any], fallback: float = 2.4) -> float:
        kg = _link_mass_kg(link)
        return min(4.2, max(1.8, (kg / max_kg) * 2.8 + 1.4)) if kg > 0 else fallback

    for link in flowsheet.get("forwardLinks") or []:
        if not isinstance(link, dict):
            continue
        src, dst = by_id.get(_clean(link.get("from"))), by_id.get(_clean(link.get("to")))
        if not src or not dst:
            continue
        points = _link_points(link, src, dst)
        known_mass = _float(link.get("massKg"), 0) > 0 or _link_mass_kg(link) > 0
        _draw_polyline(writer, points, color="172027" if known_mass else "8A949A", width=stroke_width(link), arrow=True, dash=not known_mass)
        _draw_link_label(writer, points, _link_label(link, "process"), "172027" if known_mass else "6C7680")

    for link in flowsheet.get("auxiliaryLinks") or []:
        if not isinstance(link, dict):
            continue
        src, dst = by_id.get(_clean(link.get("from"))), by_id.get(_clean(link.get("to")))
        if not src or not dst:
            continue
        kind = _clean(link.get("kind"), "waste")
        color = "657480" if kind == "vent" else "25834A" if kind == "recovery" else "965D00"
        points = _link_points(link, src, dst)
        _draw_polyline(writer, points, color=color, width=2.0, arrow=True, dash=kind in ("vent", "recovery"))
        if kind in ("waste", "vent", "recovery"):
            _draw_link_label(writer, points, _link_label(link, kind), color)

    for link in flowsheet.get("recycleLinks") or []:
        if not isinstance(link, dict):
            continue
        src, dst = by_id.get(_clean(link.get("from"))), by_id.get(_clean(link.get("to")))
        if not src or not dst:
            continue
        s = (_float(src.get("x")) + _float(src.get("w")) * 0.3, _float(src.get("y")) + _float(src.get("h")) + 12)
        e = (_float(dst.get("x")) + _float(dst.get("w")) * 0.7, _float(dst.get("y")) + _float(dst.get("h")) + 16)
        lane = max(_float(src.get("y")) + _float(src.get("h")), _float(dst.get("y")) + _float(dst.get("h"))) + 72
        fallback = [s, (s[0], lane), (e[0], lane), e]
        raw_points = _link_points(link, src, dst)
        points = raw_points if isinstance(link.get("points"), list) and len(raw_points) >= 2 else fallback
        _draw_polyline(writer, points, color="25834A", width=2.3, arrow=True, dash=True)
        _draw_link_label(writer, points, _link_label(link, "recycle") if link.get("label") else f"recycle {_clean(link.get('from'))} to {_clean(link.get('to'))}", "25834A")

    feed_box = flowsheet.get("feedBox")
    if isinstance(feed_box, dict):
        writer.rect(feed_box.get("x"), feed_box.get("y"), feed_box.get("w"), feed_box.get("h"), fill="E9F7ED", line="25834A", radius=True, name="Feed")
        writer.textbox(_float(feed_box.get("x")) + 10, _float(feed_box.get("y")) + 14, _float(feed_box.get("w")) - 20, 22, "FEED / STORAGE", size=10, color="25834A", bold=True)
        exported_feeds = flowsheet.get("feedStreams")
        feed_source = exported_feeds if isinstance(exported_feeds, list) and exported_feeds else (groups[0].get("inputStreams") if groups else [])
        feed_streams = [item for item in feed_source or [] if isinstance(item, dict)]
        for index, stream in enumerate(feed_streams[:3]):
            row_y = _float(feed_box.get("y")) + 40 + index * 22
            writer.label(
                _float(feed_box.get("x")) + 12,
                row_y - 7,
                _float(feed_box.get("w")) - 24,
                _clip_words(_stream_label(stream), 26) or "feed",
                color="172027",
                name="Feed stream",
            )

    product_box = flowsheet.get("productBox")
    if isinstance(product_box, dict):
        writer.rect(product_box.get("x"), product_box.get("y"), product_box.get("w"), product_box.get("h"), fill="E9F7ED", line="25834A", radius=True, name="Product")
        product_group_id = _clean(flowsheet.get("productGroupId"))
        product_group = by_id.get(product_group_id) or groups[-1]
        product_streams = [item for item in product_group.get("outputStreams") or [] if isinstance(item, dict)]
        product_stream = next((stream for stream in product_streams if _clean(stream.get("fate")) == "product"), product_streams[0] if product_streams else None)
        product_label = _clip_words(_clean(product_stream.get("name") if product_stream else "", "PRODUCT").upper(), 22) if product_stream else "PRODUCT"
        product_qty = _clean(product_stream.get("quantity") if product_stream else "")
        product_unit = _clean(product_stream.get("unit") if product_stream else "")
        writer.textbox(_float(product_box.get("x")) + 12, _float(product_box.get("y")) + 20, _float(product_box.get("w")) - 24, 28, product_label, size=13, color="172027", bold=True)
        if product_qty:
            writer.textbox(_float(product_box.get("x")) + 12, _float(product_box.get("y")) + 50, _float(product_box.get("w")) - 24, 18, f"{product_qty} {product_unit}".strip(), size=9.5, color="25834A", bold=True)

    for group in groups:
        fill, line = category_styles.get(_clean(group.get("category")), ("F5F9FF", "1671C2"))
        x, y, w, h = _float(group.get("x")), _float(group.get("y")), _float(group.get("w")), _float(group.get("h"))
        if group.get("concurrent"):
            writer.rect(x - 16, y - 16, w + 16, h + 16, fill="EEF2F4", line="9AA7B0", radius=True, name=f"{_clean(group.get('id'))} concurrent lane")
        writer.rect(x - 8, y - 8, w + 16, h + 16, fill="FFFFFF", line="D6E0E5", radius=True, name=f"{_clean(group.get('id'))} outer")
        writer.rect(x - 8, y - 8, 5, h + 16, fill=line, line=line, radius=True, name=f"{_clean(group.get('id'))} category stripe")
        writer.rect(x + 14, y + h * 0.52, w - 28, h * 0.40, fill=fill, line=line, radius=True, name=f"{_clean(group.get('id'))} label")
        writer.textbox(x + 22, y + 22, w - 44, 34, _category_label(group.get("category")), size=11.5, color=line, bold=True)
        writer.textbox(x + 20, y + h * 0.55, w - 40, 18, f"U{int(_float(group.get('unitNumber'), 0)) or ''} {_clean(group.get('id'))}".strip(), size=10.5, color=line, bold=True)
        writer.textbox(x + 20, y + h * 0.66, w - 40, 26, _clip_words(_clean(group.get("selectedUnit"), "unassigned unit"), 36), size=9.0, color="172027", bold=True)
        task = _clip_words(_clean(group.get("task")), 42)
        if task:
            writer.textbox(x + 20, y + h * 0.82, w - 40, 18, task, size=7.8, color="657480")
        detail_lines = [_clean(line) for line in (group.get("detailLines") or []) if _clean(line)]
        for index, line in enumerate(detail_lines[:2]):
            writer.textbox(x + 14, y + h + 22 + index * 13, w - 8, 13, _clip_words(line, 44), size=7.2, color="40515D", max_lines=1)

    _draw_stream_table(writer, flowsheet)
    _draw_title_block(writer, flowsheet, diagram_w, diagram_h)
    legend_y = _float(flowsheet.get("legendY"), 0)
    _draw_legend(writer, diagram_w, diagram_h, category_styles, legend_y if legend_y > 0 else None)

    presentation.core_properties.title = "Editable process flowsheet"
    presentation.core_properties.author = "Process Upscaling Workbench"
    presentation.core_properties.subject = "Editable process flowsheet export"
    pptx = io.BytesIO()
    presentation.save(pptx)
    return pptx.getvalue()
