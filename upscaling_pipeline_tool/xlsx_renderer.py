"""Spreadsheet export for LCI review and openLCA mapping.

This writes an .xlsx package with ordinary worksheet tables. It is intentionally
not an openLCA JSON-LD file: the workbook is the review and mapping layer between
the workbench JSON and a database-specific openLCA import.
"""

from __future__ import annotations

import io
import math
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MAX_EXCEL_ROWS = 1_048_576

DEFAULT_COLUMN_WIDTH = 18

SHEET_WIDTHS = {
    "README": [28, 82],
    "LCI Coverage": [34, 18, 78],
    "Processes": [14, 30, 20, 28, 24, 30, 16],
    "Exchanges": [14, 30, 12, 18, 14, 24, 28, 28, 12, 12, 12, 20, 24, 12, 18, 18, 14, 20, 16, 48, 42],
    "Flows": [30, 24, 22, 14, 42, 18, 22, 16],
    "Internal Links": [14, 14, 22, 38, 42],
    "OpenLCA Mapping": [16, 30, 24, 22, 16, 28, 26, 18, 48, 58, 18, 20],
    "Energy Utilities": [14, 22, 22, 18, 36, 22, 34, 48],
    "Data Quality": [20, 34, 14, 14, 58],
    "Validation": [18, 14, 18, 58, 58],
}


def _clean(value: Any, default: str = "") -> str:
    if value is None:
        return default
    text = str(value).strip()
    return re.sub(r"\s+", " ", text) if text else default


def _num(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value) if value == value else None
    text = _clean(value).replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _get(data: Any, path: str, default: Any = "") -> Any:
    cursor = data
    for part in path.split("."):
        if isinstance(cursor, dict):
            cursor = cursor.get(part)
        else:
            return default
    return default if cursor is None else cursor


def _join(values: Any) -> str:
    if not isinstance(values, list):
        return _clean(values)
    return "; ".join(_clean(value) for value in values if _clean(value))


def _excel_value(value: Any) -> Any:
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, (int, float)):
        return value if not isinstance(value, float) or math.isfinite(value) else ""
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))
    return text[:32_767]


def _sheet_name(name: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", " ", _clean(name, "Sheet")).strip()[:31] or "Sheet"
    candidate = base
    suffix = 2
    while candidate in used:
        extra = f" {suffix}"
        candidate = f"{base[:31 - len(extra)]}{extra}"
        suffix += 1
    used.add(candidate)
    return candidate


def _all_exchanges(project: dict[str, Any]) -> list[dict[str, Any]]:
    bridge = project.get("lcaBridge") if isinstance(project.get("lcaBridge"), dict) else {}
    processes = bridge.get("foregroundProcesses") if isinstance(bridge.get("foregroundProcesses"), list) else []
    exchanges: list[dict[str, Any]] = []
    for process in processes:
        if not isinstance(process, dict):
            continue
        process_id = _clean(process.get("processId"))
        process_name = _clean(process.get("processName"))
        for exchange in process.get("exchanges") or []:
            if not isinstance(exchange, dict):
                continue
            exchanges.append({**exchange, "_processId": process_id, "_processName": process_name})
    return exchanges


def _exchange_direction(exchange: dict[str, Any]) -> str:
    role = _clean(exchange.get("lcaRole")).lower()
    if role in ("technosphere_input", "internal_input", "internal_recycle_input"):
        return "input"
    if role.startswith("emission"):
        return "emission"
    if "treatment" in role or "waste" in role or role == "purge_treatment":
        return "waste output"
    return "output"


def _readme_rows(project: dict[str, Any]) -> list[list[Any]]:
    bridge = project.get("lcaBridge") if isinstance(project.get("lcaBridge"), dict) else {}
    product = bridge.get("referenceProduct") if isinstance(bridge.get("referenceProduct"), dict) else {}
    return [
        ["Field", "Value"],
        ["Workbook purpose", "LCI review and openLCA mapping handoff; not a direct openLCA JSON-LD package."],
        ["Source export", "upscaling-project.json / lcaBridge"],
        ["LCI bridge schema", bridge.get("schemaVersion", "")],
        ["LCI bridge status", bridge.get("status", "")],
        ["Reference product", product.get("rawName", "")],
        ["Reference product amount", _get(product, "amount.value")],
        ["Reference product unit", _get(product, "amount.unit")],
        ["Reference product kg", _get(product, "amount.kg")],
        ["Scale target product", _get(project, "scaleUp.basis.targetProduct")],
        ["Scale target amount", _get(project, "scaleUp.basis.targetAmount")],
        ["Scale target unit", _get(project, "scaleUp.basis.targetUnit")],
        ["OpenLCA direct import", _get(bridge, "openLcaCompatibility.directImport")],
        ["OpenLCA note", _get(bridge, "openLcaCompatibility.reason")],
        ["Next step", _get(bridge, "openLcaCompatibility.nextStep")],
    ]


def _coverage_rows(project: dict[str, Any], exchanges: list[dict[str, Any]]) -> list[list[Any]]:
    bridge = project.get("lcaBridge") if isinstance(project.get("lcaBridge"), dict) else {}
    mapping_candidates = [item for item in bridge.get("mappingCandidates") or [] if isinstance(item, dict)]
    utility_placeholders = [item for item in bridge.get("utilityPlaceholders") or [] if isinstance(item, dict)]
    rule_checks = [item for item in project.get("ruleChecks") or [] if isinstance(item, dict)]
    readiness_issues = _get(bridge, "readiness.issues", [])
    metrics = [
        ("Foreground processes", len([item for item in bridge.get("foregroundProcesses") or [] if isinstance(item, dict)]), "One worksheet row per process in the foreground model."),
        ("Total exchanges", len(exchanges), "All material exchanges collected from grouped process blocks."),
        ("Technosphere inputs", sum(1 for item in exchanges if item.get("lcaRole") == "technosphere_input"), "External product flows that need provider/database mapping."),
        ("Internal flows/recycles", sum(1 for item in exchanges if _clean(item.get("lcaRole")).startswith("internal_") or item.get("lcaRole") == "foreground_intermediate"), "Foreground links that should stay inside the model."),
        ("Emissions", sum(1 for item in exchanges if _clean(item.get("lcaRole")).startswith("emission")), "Elementary-flow candidates."),
        ("Waste/treatment outputs", sum(1 for item in exchanges if "treatment" in _clean(item.get("lcaRole")) or "waste" in _clean(item.get("lcaRole"))), "Waste flows that usually need treatment datasets."),
        ("Mapping candidates", len(mapping_candidates), "Rows requiring manual openLCA/ecoinvent dataset choice."),
        ("Utility placeholders", len(utility_placeholders), "Energy/utility rows that need duty calculation before final LCI."),
        ("Readiness issues", len(readiness_issues) if isinstance(readiness_issues, list) else 0, "Blocking or caution items from the LCI bridge."),
        ("Rule checks", len(rule_checks), "Deterministic process checks included for review context."),
    ]
    rows = [["Metric", "Value", "Interpretation"]]
    rows.extend([list(item) for item in metrics])
    return rows


def _process_rows(project: dict[str, Any], exchanges: list[dict[str, Any]]) -> list[list[Any]]:
    rows = [["process_id", "process_name", "task", "selected_unit", "blocks", "phenomena", "exchange_count"]]
    bridge = project.get("lcaBridge") if isinstance(project.get("lcaBridge"), dict) else {}
    for process in bridge.get("foregroundProcesses") or []:
        if not isinstance(process, dict):
            continue
        process_id = _clean(process.get("processId"))
        rows.append([
            process_id,
            process.get("processName", ""),
            process.get("task", ""),
            process.get("selectedUnit", ""),
            _join(process.get("blocks")),
            _join(process.get("phenomena")),
            sum(1 for exchange in exchanges if _clean(exchange.get("_processId")) == process_id),
        ])
    return rows


def _exchange_rows(exchanges: list[dict[str, Any]]) -> list[list[Any]]:
    rows = [[
        "process_id", "process_name", "block_id", "exchange_id", "direction", "lca_role",
        "raw_name", "canonical_name", "amount", "unit", "kg", "basis", "amount_status",
        "phase", "fate", "destination_group", "status", "openlca_flow_type",
        "provider_needed", "candidate_queries", "note", "source", "uncertainty_percent",
    ]]
    for exchange in exchanges:
        rows.append([
            exchange.get("_processId") or exchange.get("groupId", ""),
            exchange.get("_processName") or exchange.get("processName", ""),
            exchange.get("blockId", ""),
            exchange.get("id", ""),
            _exchange_direction(exchange),
            exchange.get("lcaRole", ""),
            exchange.get("rawName", ""),
            exchange.get("canonicalName", ""),
            _get(exchange, "amount.value"),
            _get(exchange, "amount.unit"),
            _get(exchange, "amount.kg"),
            _get(exchange, "amount.basis"),
            _get(exchange, "amount.status"),
            exchange.get("phase", ""),
            exchange.get("fate", ""),
            exchange.get("destinationGroup", ""),
            exchange.get("status", ""),
            _get(exchange, "openLcaHint.flowType"),
            _get(exchange, "openLcaHint.providerNeeded"),
            _join(exchange.get("candidateQueries")),
            exchange.get("note", ""),
            exchange.get("source", ""),
            exchange.get("uncertaintyPercent", ""),
        ])
    return rows


def _flow_rows(exchanges: list[dict[str, Any]]) -> list[list[Any]]:
    rows = [[
        "process_id", "direction", "canonical_name", "lca_role", "amount_basis",
        "flow_type_hint", "occurrences", "raw_names", "total_kg_where_available",
        "aggregation_status", "units_seen", "provider_needed",
    ]]
    grouped: dict[tuple[str, str, str, str, str], list[dict[str, Any]]] = {}
    for exchange in exchanges:
        key = (
            _clean(exchange.get("_processId") or exchange.get("groupId")),
            _exchange_direction(exchange),
            _clean(exchange.get("canonicalName") or exchange.get("rawName")),
            _clean(exchange.get("lcaRole")),
            _clean(_get(exchange, "amount.basis")),
        )
        grouped.setdefault(key, []).append(exchange)
    for (process_id, direction, name, role, basis), items in sorted(grouped.items()):
        raw_names = sorted({_clean(item.get("rawName")) for item in items if _clean(item.get("rawName"))})
        units = sorted({_clean(_get(item, "amount.unit")) for item in items if _clean(_get(item, "amount.unit"))})
        kg_values = [_num(_get(item, "amount.kg")) for item in items]
        complete_kg = all(value is not None for value in kg_values)
        total_kg = sum(value for value in kg_values if value is not None) if complete_kg else None
        aggregation_status = "same process/direction/basis; kg-convertible" if complete_kg else "not summed: one or more amounts are not kg-convertible"
        rows.append([
            process_id,
            direction,
            name,
            role,
            basis,
            _get(items[0], "openLcaHint.flowType"),
            len(items),
            "; ".join(raw_names),
            total_kg if total_kg is not None else "",
            aggregation_status,
            "; ".join(units),
            _get(items[0], "openLcaHint.providerNeeded"),
        ])
    return rows


def _link_rows(project: dict[str, Any], exchanges: list[dict[str, Any]]) -> list[list[Any]]:
    rows = [["from", "to", "source", "related_exchange_ids", "related_stream_names"]]
    for link in project.get("links") or []:
        if not isinstance(link, dict):
            continue
        from_id, to_id = _clean(link.get("from")), _clean(link.get("to"))
        related = [exchange for exchange in exchanges if _clean(exchange.get("groupId")) == from_id and _clean(exchange.get("destinationGroup")).upper() == to_id.upper()]
        rows.append([
            from_id,
            to_id,
            "declared process link",
            _join([item.get("id", "") for item in related]),
            _join([item.get("rawName", "") for item in related]),
        ])
    return rows


def _mapping_rows(project: dict[str, Any]) -> list[list[Any]]:
    rows = [[
        "mapping_status", "canonical_name", "lca_role", "openlca_flow_type_hint",
        "provider_needed", "selected_openlca_flow_id", "selected_provider_id",
        "selected_location", "candidate_queries", "occurrences", "total_kg_where_available",
        "aggregation_status", "units_seen",
    ]]
    bridge = project.get("lcaBridge") if isinstance(project.get("lcaBridge"), dict) else {}
    for candidate in bridge.get("mappingCandidates") or []:
        if not isinstance(candidate, dict):
            continue
        occurrences = []
        for occurrence in candidate.get("occurrences") or []:
            if isinstance(occurrence, dict):
                occurrences.append(f"{_clean(occurrence.get('groupId'))}/{_clean(occurrence.get('blockId'))}: {_clean(occurrence.get('rawName'))}")
        amount_items = [occurrence for occurrence in candidate.get("occurrences") or [] if isinstance(occurrence, dict)]
        kg_values = [_num(_get(occurrence, "amount.kg")) for occurrence in amount_items]
        bases = {_clean(_get(occurrence, "amount.basis")) for occurrence in amount_items if _clean(_get(occurrence, "amount.basis"))}
        can_sum = bool(amount_items) and all(value is not None for value in kg_values) and len(bases) <= 1
        total_kg = sum(value for value in kg_values if value is not None) if can_sum else None
        aggregation_status = "compatible basis; kg-convertible" if can_sum else "not summed: mixed basis or non-kg amount"
        units_seen = sorted({_clean(_get(occurrence, "amount.unit")) for occurrence in candidate.get("occurrences") or [] if isinstance(occurrence, dict) and _clean(_get(occurrence, "amount.unit"))})
        rows.append([
            candidate.get("mappingStatus", "unmapped"),
            candidate.get("canonicalName", ""),
            candidate.get("lcaRole", ""),
            _get(candidate, "openLcaHint.flowType"),
            _get(candidate, "openLcaHint.providerNeeded"),
            candidate.get("selectedOpenLcaFlowId", ""),
            candidate.get("selectedProviderId", ""),
            candidate.get("selectedLocation", ""),
            _join(candidate.get("candidateQueries")),
            "; ".join(occurrences),
            total_kg if total_kg is not None else "",
            aggregation_status,
            "; ".join(units_seen),
        ])
    return rows


def _energy_rows(project: dict[str, Any]) -> list[list[Any]]:
    rows = [["group_id", "task", "event_type", "data_status", "missing", "lca_role", "mapping_status", "candidate_queries"]]
    bridge = project.get("lcaBridge") if isinstance(project.get("lcaBridge"), dict) else {}
    for item in bridge.get("utilityPlaceholders") or []:
        if not isinstance(item, dict):
            continue
        rows.append([
            item.get("groupId", ""),
            item.get("task", ""),
            item.get("eventType", ""),
            item.get("dataStatus", ""),
            _join(item.get("missing")),
            item.get("lcaRole", ""),
            item.get("mappingStatus", ""),
            _join(item.get("candidateQueries")),
        ])
    return rows


def _data_quality_rows(project: dict[str, Any]) -> list[list[Any]]:
    rows = [["category", "item", "level", "status", "note"]]
    readiness = project.get("dataReadiness") if isinstance(project.get("dataReadiness"), dict) else {}
    for category in readiness.get("categories") or []:
        if not isinstance(category, dict):
            continue
        for item in category.get("items") or []:
            if not isinstance(item, dict):
                continue
            status = "ok" if item.get("ok") is True else "missing" if item.get("ok") is False else "confirm"
            rows.append([category.get("name", ""), item.get("name", ""), item.get("level", ""), status, item.get("note", "")])
    return rows


def _validation_rows(project: dict[str, Any]) -> list[list[Any]]:
    rows = [["source", "severity", "target", "issue", "recommendation"]]
    bridge = project.get("lcaBridge") if isinstance(project.get("lcaBridge"), dict) else {}
    for issue in _get(bridge, "readiness.issues", []):
        rows.append(["lcaBridge", "", "", issue, "Resolve before openLCA JSON-LD generation."])
    for issue in project.get("ruleChecks") or []:
        if not isinstance(issue, dict):
            continue
        rows.append(["ruleChecks", issue.get("severity", ""), issue.get("target", ""), issue.get("title", issue.get("issue", "")), issue.get("recommendation", "")])
    return rows


def render_lci_workbook_xlsx(project: dict[str, Any]) -> bytes:
    if not isinstance(project, dict):
        raise ValueError("Missing project export payload.")

    exchanges = _all_exchanges(project)
    used_names: set[str] = set()
    sheets = [
        (_sheet_name("README", used_names), _readme_rows(project)),
        (_sheet_name("LCI Coverage", used_names), _coverage_rows(project, exchanges)),
        (_sheet_name("Processes", used_names), _process_rows(project, exchanges)),
        (_sheet_name("Exchanges", used_names), _exchange_rows(exchanges)),
        (_sheet_name("Flows", used_names), _flow_rows(exchanges)),
        (_sheet_name("Internal Links", used_names), _link_rows(project, exchanges)),
        (_sheet_name("OpenLCA Mapping", used_names), _mapping_rows(project)),
        (_sheet_name("Energy Utilities", used_names), _energy_rows(project)),
        (_sheet_name("Data Quality", used_names), _data_quality_rows(project)),
        (_sheet_name("Validation", used_names), _validation_rows(project)),
    ]

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "Process Upscaling Workbench"
    workbook.properties.title = "LCI review workbook"
    workbook.properties.description = "Editable LCI review and openLCA mapping handoff"

    edge = Side(style="thin", color="D9E2EA")
    border = Border(left=edge, right=edge, top=edge, bottom=edge)
    header_fill = PatternFill("solid", fgColor="365F7E")
    header_font = Font(name="Aptos", size=11, bold=True, color="FFFFFF")
    body_font = Font(name="Aptos", size=10, color="172027")

    for name, rows in sheets:
        worksheet = workbook.create_sheet(name)
        safe_rows = rows[:MAX_EXCEL_ROWS]
        for row in safe_rows:
            worksheet.append([_excel_value(value) for value in row])

        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        if len(safe_rows) > 1:
            worksheet.auto_filter.ref = worksheet.dimensions

        widths = SHEET_WIDTHS.get(name, [])
        for index in range(1, worksheet.max_column + 1):
            worksheet.column_dimensions[get_column_letter(index)].width = (
                widths[index - 1] if index <= len(widths) else DEFAULT_COLUMN_WIDTH
            )

        worksheet.row_dimensions[1].height = 30
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = border

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
