#!/usr/bin/env python3
"""Regression check for editable PowerPoint flowsheet export."""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from upscaling_pipeline_tool.pptx_renderer import render_flowsheet_pptx


payload = {
    "flowsheet": {
        "title": "Generated Process Flowsheet",
        "basis": "test basis",
        "width": 1000,
        "height": 620,
        "productGroupId": "G2",
        "feedBox": {"x": 30, "y": 120, "w": 150, "h": 120},
        "productBox": {"x": 820, "y": 140, "w": 140, "h": 90},
        "dischargeBox": {
            "x": 820,
            "y": 266,
            "w": 190,
            "h": 56,
            "rows": [{"id": "wastewater", "label": "to wastewater treatment", "kind": "waste", "kg": 1.5, "unknown": 0, "count": 1, "units": ["G1"]}],
        },
        "groups": [
            {
                "id": "G1",
                "unitNumber": 1,
                "task": "reaction and hold under nitrogen",
                "selectedUnit": "stirred tank reactor with condenser",
                "category": "reactor",
                "subcategory": "reactor",
                "x": 220,
                "y": 120,
                "w": 190,
                "h": 160,
                "totalOutputKg": 10,
                "inputStreams": [{"name": "ethyl acetate solvent", "quantity": "12", "unit": "kg"}],
                "outputStreams": [{"name": "reaction slurry intermediate", "quantity": "8", "unit": "kg"}],
                "layout": {
                    "accent": "#CF4B42",
                    "tagTop": 250,
                    "tagHeight": 44,
                    "unitLines": ["stirred tank reactor"],
                    "footerLine": "10 kg/batch main route",
                    "taskLine": "reaction and hold under nitrogen",
                    "taskY": 314,
                    "productY": 330,
                },
                "symbol": {
                    "x": 210,
                    "y": 110,
                    "w": 210,
                    "h": 130,
                    "png": "data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==",
                    "svg": "<svg xmlns='http://www.w3.org/2000/svg' width='210' height='130'><rect x='4' y='4' width='200' height='120' fill='none' stroke='#CF4B42'/></svg>",
                },
                "boundaryOutlets": [
                    {
                        "id": "wastewater",
                        "short": "to WWT",
                        "label": "to wastewater treatment",
                        "kind": "waste",
                        "tag": "S9",
                        "kg": 1.5,
                        "unknown": 0,
                        "streams": [{"name": "spent wash water", "quantity": "1.5", "unit": "kg"}],
                        "points": [{"x": 364, "y": 280}, {"x": 364, "y": 338}],
                    }
                ],
            },
            {
                "id": "G2",
                "unitNumber": 2,
                "task": "filtration and product isolation",
                "selectedUnit": "pressure nutsche filter dryer",
                "category": "separation",
                "subcategory": "filter",
                "x": 560,
                "y": 120,
                "w": 190,
                "h": 160,
                "totalOutputKg": 8,
                "outputStreams": [{"name": "final product", "quantity": "5.2", "unit": "kg", "fate": "product"}],
            },
        ],
        "forwardLinks": [
            {"from": "G1", "to": "G2", "directStreams": [{"name": "reaction slurry intermediate", "quantity": "8", "unit": "kg"}]}
        ],
        "auxiliaryLinks": [
            {"from": "G1", "to": "G2", "kind": "waste", "directStreams": [{"name": "mother liquor", "quantity": "2", "unit": "kg"}]}
        ],
        "recycleLinks": [
            {"from": "G2", "to": "G1", "directStreams": [{"name": "recovered solvent", "quantity": "1", "unit": "kg"}]}
        ],
    }
}


body = render_flowsheet_pptx(payload)
assert body.startswith(b"PK"), "PowerPoint export should be an OOXML zip package"
deck = Presentation(io.BytesIO(body))
assert len(deck.slides) == 1, "PowerPoint export should reopen with one slide"
assert len(deck.slides[0].shapes) > 10, "PowerPoint export should retain editable native shapes"
roundtrip = io.BytesIO()
deck.save(roundtrip)
assert len(Presentation(io.BytesIO(roundtrip.getvalue())).slides) == 1, "PowerPoint export should survive a save round trip"
with zipfile.ZipFile(io.BytesIO(body)) as zf:
    slide = zf.read("ppt/slides/slide1.xml").decode("utf-8")

assert "Legend Panel" in slide, "PowerPoint export should include an editable legend panel"
assert "Process" in slide and "Recycle" in slide and "Vent/VOC" in slide, "Legend should include line meanings"
assert "Reactor" in slide and "Separation" in slide, "Legend should include unit categories"
assert '<a:tailEnd type="triangle"/>' in slide, "Connector arrowheads should point at the line endpoint"
assert "reaction..." in slide or "reaction slurry" in slide, "Process stream label should be included as editable text"
assert "S9 to WWT" in slide, "Boundary outlets should be exported as tagged off-page connectors"

# Equipment symbols: PowerPoint has no jacketed-reactor shape, so the symbol travels as a picture
# with the vector attached. Each unit is one group, so the whole card moves together.
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Picture

shapes = list(deck.slides[0].shapes)
groups = [shape for shape in shapes if isinstance(shape, GroupShape)]
unit_groups = [group for group in groups if group.name.startswith("U1 G1")]
assert unit_groups, "Each unit should be exported as one named group, got " + str([s.name for s in shapes])
unit = unit_groups[0]
pictures = [child for child in unit.shapes if isinstance(child, Picture)]
assert pictures, "The unit group should contain its equipment symbol as a picture"
frame = next(child for child in unit.shapes if child.name.endswith("frame"))
picture = pictures[0]
assert picture.left >= frame.left and picture.left + picture.width <= frame.left + frame.width, "The symbol picture must sit inside the unit card, not overhang it"
assert [child.name for child in unit.shapes].index(picture.name) < [child.name for child in unit.shapes].index("G1 tag"), "The symbol must sit behind the unit's label, as a background"
assert any("G1 to G2 (process)" in group.name for group in groups), "Each stream should be one named group carrying its arrow and label"

with zipfile.ZipFile(io.BytesIO(body)) as package:
    names = package.namelist()
    assert any(name.endswith(".svg") for name in names), "The vector symbol should be stored in the package"
    slide_xml = package.read("ppt/slides/slide1.xml").decode()
    assert "svgBlip" in slide_xml, "The picture should reference its vector so PowerPoint renders the SVG"
    assert "image/svg+xml" in package.read("[Content_Types].xml").decode(), "The SVG part needs a registered content type or PowerPoint repairs the file"

# An SVG carrying a script is not embedded; the PNG still is.
import copy

unsafe = copy.deepcopy(payload)
unsafe["flowsheet"]["groups"][0]["symbol"]["svg"] = "<svg xmlns='http://www.w3.org/2000/svg'><script>alert(1)</script></svg>"
unsafe_body = render_flowsheet_pptx(unsafe)
with zipfile.ZipFile(io.BytesIO(unsafe_body)) as package:
    assert not any(name.endswith(".svg") for name in package.namelist()), "An SVG containing a script must not be embedded"
    assert any(name.endswith(".png") for name in package.namelist()), "The PNG fallback should still be exported"
assert "DISCHARGES" in slide and "wastewater treatment" in slide, "The discharge box should be exported with its per-boundary totals"

print("Flowsheet PowerPoint export regression check passed.")
